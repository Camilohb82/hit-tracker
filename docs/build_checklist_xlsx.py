# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

OUT = "/home/user/hit-tracker/docs/Checklist-CEDI-Madrid.xlsx"

# ---------- paleta ----------
NAVY   = "0E4E7C"
NAVY_D = "0A3A5C"
GREY_H = "E8ECEF"
GREY_L = "F5F7F9"
YELLOW = "FFF2CC"
TAPE   = "B98400"
WHITE  = "FFFFFF"

F  = "Arial"
def font(sz=10, b=False, color="000000", it=False):
    return Font(name=F, size=sz, bold=b, color=color, italic=it)

thin = Side(style="thin", color="BFC9D1")
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)

FILL_H  = PatternFill("solid", fgColor=NAVY)
FILL_S  = PatternFill("solid", fgColor=GREY_H)
FILL_IN = PatternFill("solid", fgColor=YELLOW)
FILL_C  = PatternFill("solid", fgColor=GREY_L)

WRAP  = Alignment(wrap_text=True, vertical="top")
WRAPC = Alignment(wrap_text=True, vertical="center")
CTR   = Alignment(horizontal="center", vertical="center")
CTRW  = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

def title_block(ws, title, subtitle, ncols):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    c = ws.cell(row=1, column=1, value=title)
    c.font = font(14, True, NAVY_D); c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
    c = ws.cell(row=2, column=2 - 1, value=subtitle)
    c.font = font(9, False, "47535F"); c.alignment = WRAPC
    ws.row_dimensions[2].height = 30

def headers(ws, row, labels, widths):
    for i, (lab, w) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=i, value=lab)
        c.font = font(9, True, WHITE); c.fill = FILL_H
        c.alignment = CTRW; c.border = BOX
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30

def style_range(ws, r1, r2, c1, c2, inputs=(), calcs=()):
    for r in range(r1, r2 + 1):
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.font = font(10)
            cell.alignment = WRAP
            if c in inputs:
                cell.fill = FILL_IN; cell.font = font(10, color="0000FF")
            elif c in calcs:
                cell.fill = FILL_C

def example_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = font(9, color="8A8F96", it=True)
        cell.fill = PatternFill("solid", fgColor="FAFBFC")
        cell.border = BOX
        cell.alignment = WRAP

def si_no(ws, rng):
    dv = DataValidation(type="list", formula1='"Sí,No"', allow_blank=True)
    dv.error = "Selecciona Sí o No"; dv.errorTitle = "Valor no válido"
    ws.add_data_validation(dv); dv.add(rng)

# =========================================================
# PORTADA
# =========================================================
ws = wb.active; ws.title = "Portada"
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [3, 34, 40, 22, 3]):
    ws.column_dimensions[col].width = w

ws.merge_cells("B2:D2")
c = ws["B2"]; c.value = "EVALUACIÓN OPERATIVA — CEDI MADRID"
c.font = font(18, True, NAVY_D)
ws.row_dimensions[2].height = 30

ws.merge_cells("B3:D3")
c = ws["B3"]; c.value = "Formato de captura en piso — Día 1"
c.font = font(11, False, TAPE)

ws.merge_cells("B5:D6")
c = ws["B5"]
c.value = ("Hoy no se diagnostica: se captura evidencia propia que después nadie pueda discutir. "
           "Diligencia solo las celdas amarillas; las grises se calculan solas. "
           "La hoja Resumen se llena automáticamente a medida que avanzas.")
c.font = font(10, color="47535F"); c.alignment = WRAP
ws.row_dimensions[5].height = 16

ws.merge_cells("B8:D8")
c = ws["B8"]; c.value = "DATOS DE LA VISITA"
c.font = font(10, True, WHITE); c.fill = FILL_H; c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[8].height = 20

visita = [
    ("Centro de distribución", "CEDI Madrid"),
    ("Fecha de la visita", ""),
    ("Responsable de la evaluación", ""),
    ("Jefe del CEDI", ""),
    ("Turno observado", ""),
    ("Hora de corte de pedidos", ""),
    ("Hora de salida del último camión", ""),
    ("Día de la semana", ""),
]
r = 9
for lab, val in visita:
    ws.cell(row=r, column=2, value=lab).font = font(10, True)
    ws.cell(row=r, column=2).border = BOX
    ws.cell(row=r, column=2).fill = FILL_S
    cc = ws.cell(row=r, column=3, value=val)
    cc.fill = FILL_IN; cc.font = font(10, color="0000FF"); cc.border = BOX
    ws.cell(row=r, column=4).border = BOX
    ws.row_dimensions[r].height = 18
    r += 1

ws.merge_cells("B19:D19")
c = ws["B19"]; c.value = "CÓMO SE USA"
c.font = font(10, True, WHITE); c.fill = FILL_H; c.alignment = Alignment(vertical="center", indent=1)
ws.row_dimensions[19].height = 20

guia = [
    ("Amarillo", "Celdas que diligencias tú en piso."),
    ("Gris claro", "Se calculan automáticamente. No las escribas."),
    ("Fila EJEMPLO", "Muestra el formato esperado. Déjala o bórrala, no entra en los totales."),
    ("Orden sugerido", "1 Checklist → arranca los relojes (etiquetas, portería, conteo ciego) → recorre a contracorriente → captura en las hojas 2 a 8 → cierra con 9 y 10."),
    ("Resumen", "Se actualiza solo. Es lo que muestras al cierre del día."),
]
r = 20
for lab, txt in guia:
    cc = ws.cell(row=r, column=2, value=lab); cc.font = font(10, True); cc.border = BOX; cc.fill = FILL_S
    cc = ws.cell(row=r, column=3, value=txt); cc.font = font(10); cc.border = BOX; cc.alignment = WRAP
    ws.cell(row=r, column=4).border = BOX
    ws.row_dimensions[r].height = 28
    r += 1

ws.merge_cells("B27:D29")
c = ws["B27"]
c.value = ("Nota sobre los valores de referencia: los cortes usados en este archivo (85–90% de ocupación, "
           "50% de desplazamiento en el ciclo, 95% de ERI, 90 minutos de permanencia en muelle) son referencias "
           "de industria para CEDI de consumo masivo con WMS y radiofrecuencia. Sirven para leer lo que midas hoy; "
           "conviértelos en meta solo después de calibrarlos contra el histórico propio de Madrid.")
c.font = font(9, color="74818D"); c.alignment = WRAP

# =========================================================
# RESUMEN
# =========================================================
rs = wb.create_sheet("Resumen")
rs.sheet_view.showGridLines = False
title_block(rs, "RESUMEN DEL DÍA", "Se llena automáticamente desde las hojas de captura. Es la evidencia que devuelves en la reunión de cierre.", 4)
headers(rs, 4, ["Indicador", "Resultado", "Referencia", "Lectura"], [42, 14, 16, 62])
rs.freeze_panes = "A5"

# (label, formula_origen, num_format, referencia, expr_lectura)
res = [
    ("Ocupación de posiciones", "'2 Ocupación'!$F$16", "0.0%", "80% – 85%",
     'IF($B5="","Pendiente de medir",IF($B5>0.9,"CRÍTICO. Zona de congestión: la productividad que midas hoy está contaminada.",IF($B5>0.85,"ATENCIÓN. En el límite; cada punto adicional ya cuesta caro.","OK. Si igual se siente lleno, el problema es de flujo, no de espacio.")))'),
    ("Capacidad fantasma (posiciones parciales)", "'2 Ocupación'!$G$16", "0.0%", "< 10%",
     'IF($B6="","Pendiente de medir",IF($B6>0.15,"CRÍTICO. Posiciones que el sistema ve ocupadas y no almacenan nada.",IF($B6>0.1,"ATENCIÓN. Consolidar parciales libera espacio sin comprar nada.","OK. Sin efecto panal relevante.")))'),
    ("Desplazamiento sobre el ciclo de alistamiento", "'3 Ciclo picking'!$F$14", "0.0%", "< 40%",
     'IF($B7="","Pendiente de medir",IF($B7>0.5,"CRÍTICO. El problema es el slotting, no la gente. Re-slotting antes que metas de velocidad.",IF($B7>0.4,"ATENCIÓN. Vale un re-slotting de los SKU de mayor rotación.","OK. Busca la pérdida en búsqueda, espera o reproceso.")))'),
    ("ERI por ubicación (conteo ciego)", "'4 Conteo ciego'!$F$4", "0.0%", "> 97%",
     'IF($B8="","Pendiente de medir",IF($B8<0.95,"CRÍTICO. Prioridad uno: exactitud de inventario antes de slotting, método o chequeo.",IF($B8<0.97,"ATENCIÓN. Subir frecuencia de conteos cíclicos en zonas A.","OK. Los errores de despacho tienen otra causa; búscala en el Pareto.")))'),
    ("Tiempo que no agrega valor (muestreo)", "'5 Muestreo'!$H$14", "0.0%", "< 35%",
     'IF($B9="","Pendiente de medir",IF($B9>0.5,"CRÍTICO. Más de la mitad del día no agrega valor. Es diseño, no esfuerzo.",IF($B9>0.35,"ATENCIÓN. Rango habitual en CEDI sin estándares de trabajo.","OK. Operación con poco desperdicio visible.")))'),
    ("Permanencia promedio de vehículo en muelle (min)", "'6 Vehículos'!$H$26", "0", "< 90 min",
     'IF($B10="","Pendiente de medir",IF($B10>180,"CRÍTICO. El cuello de botella está en muelle o staging, no en alistamiento.",IF($B10>90,"ATENCIÓN. Revisar programación de citas y secuencia de cargue.","OK. Muelle fluido.")))'),
    ("Pallets en piso etiquetados hoy", "'7 Pallets en piso'!$B$16", "0", "0",
     'IF($B11="","Pendiente de medir",IF($B11>30,"CRÍTICO. Backlog físico grande. Verificar a las 48 h cuántos siguen ahí.",IF($B11>0,"ATENCIÓN. Verificar a las 48 h cuáles no se movieron.","OK. Sin producto fuera de posición.")))'),
    ("Agotados internos de picking en el turno", "'8 Agotados'!$B$46", "0", "0",
     'IF($B12="","Pendiente de medir",IF($B12>10,"CRÍTICO. El reabastecimiento va detrás del picking. Pasar a reabastecimiento en ola.",IF($B12>0,"ATENCIÓN. Cuantificar minutos perdidos por evento.","OK. Reabastecimiento va adelante del picking.")))'),
    ("Avance del checklist del día", "'1 Checklist'!$D$3", "0.0%", "100%",
     'IF($B13="","Sin iniciar",IF($B13<1,"Quedan acciones pendientes antes de cerrar el día.","Checklist completo."))'),
    ("Solicitud de datos entregada y recibida", "'10 Solicitud datos'!$G$3", "0.0%", "100%",
     'IF($B14="","Sin entregar",IF($B14<1,"Pendiente: sin el dato crudo no hay análisis la próxima semana.","Información completa recibida."))'),
]
r = 5
for lab, src, nf, ref, lect in res:
    rs.cell(row=r, column=1, value=lab).font = font(10, True)
    b = rs.cell(row=r, column=2, value='=IF({0}="","",{0})'.format(src))
    b.number_format = nf; b.alignment = CTR; b.fill = FILL_C; b.font = font(11, True)
    cc = rs.cell(row=r, column=3, value=ref); cc.alignment = CTR; cc.font = font(9, color="74818D")
    d = rs.cell(row=r, column=4, value="=" + lect); d.font = font(9, color="47535F"); d.alignment = WRAP
    for cnum in range(1, 5):
        rs.cell(row=r, column=cnum).border = BOX
    rs.row_dimensions[r].height = 32
    r += 1
for c1 in range(1, 5):
    rs.cell(row=4, column=c1).alignment = CTRW
rs.cell(row=5, column=1).alignment = WRAPC

rs.merge_cells("A17:D19")
c = rs["A17"]
c.value = ("Con estas evidencias ya no dependes de que nadie te cuente qué pasa en Madrid: ocupación medida por ti, "
           "desplazamiento cronometrado, ERI de conteo ciego, muestreo de trabajo, pallets en piso con fecha puesta "
           "y la solicitud de datos entregada. En la reunión de cierre devuelve hechos con número, no opiniones.")
c.font = font(9, color="74818D"); c.alignment = WRAP

# =========================================================
# 1 CHECKLIST
# =========================================================
ck = wb.create_sheet("1 Checklist")
ck.sheet_view.showGridLines = False
title_block(ck, "CHECKLIST DEL DÍA", "Marca Sí en la columna Hecho. El avance se calcula solo y se muestra en Resumen.", 7)
ck["C3"] = "Avance del checklist"
ck["C3"].font = font(10, True); ck["C3"].alignment = Alignment(horizontal="right")
ck["D3"] = '=IF(COUNTA($C$6:$C$25)=0,"",COUNTIF($D$6:$D$25,"Sí")/COUNTA($C$6:$C$25))'
ck["D3"].number_format = "0.0%"; ck["D3"].font = font(11, True, NAVY_D)
ck["D3"].fill = FILL_C; ck["D3"].alignment = CTR; ck["D3"].border = BOX
headers(ck, 5, ["#", "Bloque", "Acción", "Hecho", "Hora", "Responsable", "Observación"],
        [5, 20, 60, 9, 9, 18, 34])
ck.freeze_panes = "A6"

acciones = [
    ("Antes de bajarte", "Reunión de apertura de 10 minutos, de pie. Nada de sala de juntas."),
    ("Antes de bajarte", "Preguntar hora de corte de pedidos y hora del último camión."),
    ("Antes de bajarte", "Anunciar en el primer minuto: se mide el proceso, no a las personas."),
    ("Antes de bajarte", "Verificar EPP: botas, chaleco, casco si aplica."),
    ("Primeros 30 min", "Etiquetar con fecha y hora cada pallet que esté en piso. Foto del conjunto."),
    ("Primeros 30 min", "Abrir planilla en portería: placa, llegada, entrada a muelle, salida."),
    ("Primeros 30 min", "Encargar conteo ciego de 100 ubicaciones elegidas por mí, no por ellos."),
    ("Primeros 30 min", "Seis fotos con hora de puntos fijos (repetir en pico y al cierre)."),
    ("Recorrido", "Recorrer a contracorriente: del muelle de salida hacia atrás hasta recepción."),
    ("Recorrido", "Repetir el recorrido en el pico, antes del corte de pedidos."),
    ("Mediciones", "Hoja 2 — Ocupación: contar 5 pasillos al azar."),
    ("Mediciones", "Hoja 3 — Ciclo de alistamiento: cronometrar 4 operarios, ciclo completo."),
    ("Mediciones", "Hoja 4 — Conteo ciego: registrar resultado y calcular ERI."),
    ("Mediciones", "Hoja 5 — Muestreo de trabajo: una ronda cada 20 minutos."),
    ("Mediciones", "Hoja 6 — Vehículos: recoger la planilla de portería."),
    ("Mediciones", "Hoja 7 — Pallets en piso: contar y fechar por zona."),
    ("Mediciones", "Hoja 8 — Agotados de picking: dejar el registro con un supervisor."),
    ("Cierre", "Devolver tres hechos con número al jefe del CEDI. Sin opiniones."),
    ("Cierre", "Entregar la solicitud de datos firmada, con responsable y fecha."),
    ("Cierre", "Acordar una sola acción que empieza mañana y fijar la segunda visita."),
]
r = 6
for i, (bloque, acc) in enumerate(acciones, start=1):
    ck.cell(row=r, column=1, value=i)
    ck.cell(row=r, column=2, value=bloque)
    ck.cell(row=r, column=3, value=acc)
    ck.row_dimensions[r].height = 26
    r += 1
style_range(ck, 6, 25, 1, 7, inputs=(4, 5, 6, 7))
for rr in range(6, 26):
    ck.cell(row=rr, column=1).alignment = CTR
    ck.cell(row=rr, column=2).font = font(9, True, TAPE)
    ck.cell(row=rr, column=4).alignment = CTR
si_no(ck, "D6:D25")
ck.conditional_formatting.add("D6:D25",
    CellIsRule(operator="equal", formula=['"Sí"'], fill=PatternFill("solid", fgColor="DFEFE7")))

# =========================================================
# 2 OCUPACIÓN
# =========================================================
oc = wb.create_sheet("2 Ocupación")
oc.sheet_view.showGridLines = False
title_block(oc, "OCUPACIÓN REAL DE POSICIONES",
            "Elige 5 pasillos al azar — al azar de verdad, no los que te muestren. Parcial = una caja o medio pallet bloqueando una posición entera: el sistema la ve ocupada y no almacena nada.", 7)
headers(oc, 4, ["Pasillo / zona", "Ocupadas", "Parciales", "Vacías", "Total posiciones",
                "% Ocupación", "% Capacidad fantasma"], [24, 12, 12, 12, 15, 13, 18])
oc.freeze_panes = "A5"

oc.cell(row=5, column=1, value="EJEMPLO — Pasillo 12")
oc.cell(row=5, column=2, value=78); oc.cell(row=5, column=3, value=14); oc.cell(row=5, column=4, value=8)
oc.cell(row=5, column=5, value='=IF(SUM(B5:D5)=0,"",SUM(B5:D5))')
oc.cell(row=5, column=6, value='=IF($E5="","",($B5+$C5)/$E5)')
oc.cell(row=5, column=7, value='=IF($E5="","",$C5/$E5)')
example_row(oc, 5, 7)
oc.cell(row=5, column=6).number_format = "0.0%"
oc.cell(row=5, column=7).number_format = "0.0%"

for r in range(6, 16):
    oc.cell(row=r, column=5, value='=IF(SUM(B{0}:D{0})=0,"",SUM(B{0}:D{0}))'.format(r))
    oc.cell(row=r, column=6, value='=IF($E{0}="","",($B{0}+$C{0})/$E{0})'.format(r))
    oc.cell(row=r, column=7, value='=IF($E{0}="","",$C{0}/$E{0})'.format(r))
    oc.row_dimensions[r].height = 18
style_range(oc, 6, 15, 1, 7, inputs=(1, 2, 3, 4), calcs=(5, 6, 7))
for r in range(6, 16):
    for c in (2, 3, 4, 5, 6, 7):
        oc.cell(row=r, column=c).alignment = CTR
    oc.cell(row=r, column=6).number_format = "0.0%"
    oc.cell(row=r, column=7).number_format = "0.0%"

oc.cell(row=16, column=1, value="TOTAL MUESTRA")
for c, f_ in [(2, "=SUM(B6:B15)"), (3, "=SUM(C6:C15)"), (4, "=SUM(D6:D15)"), (5, "=SUM(E6:E15)")]:
    oc.cell(row=16, column=c, value=f_)
oc.cell(row=16, column=6, value='=IF($E$16=0,"",($B$16+$C$16)/$E$16)')
oc.cell(row=16, column=7, value='=IF($E$16=0,"",$C$16/$E$16)')
for c in range(1, 8):
    cell = oc.cell(row=16, column=c)
    cell.font = font(11, True); cell.fill = FILL_S; cell.border = BOX
    cell.alignment = CTR if c > 1 else Alignment(vertical="center")
oc.cell(row=16, column=6).number_format = "0.0%"
oc.cell(row=16, column=7).number_format = "0.0%"
oc.row_dimensions[16].height = 22
oc.conditional_formatting.add("F16",
    CellIsRule(operator="greaterThan", formula=["0.9"], fill=PatternFill("solid", fgColor="F9E4E3"), font=font(11, True, "A32320")))
oc.conditional_formatting.add("F16",
    CellIsRule(operator="between", formula=["0.85", "0.9"], fill=PatternFill("solid", fgColor="F8EDD8")))

oc.merge_cells("A18:G19")
c = oc["A18"]
c.value = ("Lectura: por encima de 85% cada punto adicional de ocupación cuesta cada vez más, y por encima de 90% "
           "la operación entra en congestión — más doble manipulación, más búsqueda, más errores. Mientras estés ahí, "
           "cualquier medición de productividad que tomes está contaminada.")
c.font = font(9, color="74818D"); c.alignment = WRAP

# =========================================================
# 3 CICLO PICKING
# =========================================================
cp = wb.create_sheet("3 Ciclo picking")
cp.sheet_view.showGridLines = False
title_block(cp, "CICLO DE ALISTAMIENTO",
            "Sigue a 4 operarios distintos, un ciclo completo cada uno. Cronómetro con dos tiempos separados: caminando y tomando producto (incluye chequear y rotular).", 13)
headers(cp, 4, ["Operario", "Zona", "Seg. caminando", "Seg. tomando", "Ciclo total (s)",
                "% desplazamiento", "Líneas del ciclo", "Seg. por línea",
                "No encontró", "Posición vacía", "Devolvió", "Preguntó", "¿Escaneó cada línea?"],
        [16, 12, 13, 13, 13, 14, 12, 12, 11, 11, 10, 10, 16])
cp.freeze_panes = "C5"

ej = ["EJEMPLO — J. Ramírez", "Picking A", 412, 298, None, None, 22, None, 3, 1, 0, 2, "No"]
for i, v in enumerate(ej, start=1):
    cp.cell(row=5, column=i, value=v)
cp.cell(row=5, column=5, value='=IF(OR($C5="",$D5=""),"",$C5+$D5)')
cp.cell(row=5, column=6, value='=IF($E5="","",$C5/$E5)')
cp.cell(row=5, column=8, value='=IF(OR($E5="",$G5=""),"",IFERROR($E5/$G5,""))')
example_row(cp, 5, 13)
cp.cell(row=5, column=6).number_format = "0.0%"
cp.cell(row=5, column=8).number_format = "0.0"

for r in range(6, 14):
    cp.cell(row=r, column=5, value='=IF(OR($C{0}="",$D{0}=""),"",$C{0}+$D{0})'.format(r))
    cp.cell(row=r, column=6, value='=IF($E{0}="","",$C{0}/$E{0})'.format(r))
    cp.cell(row=r, column=8, value='=IF(OR($E{0}="",$G{0}=""),"",IFERROR($E{0}/$G{0},""))'.format(r))
    cp.row_dimensions[r].height = 18
style_range(cp, 6, 13, 1, 13, inputs=(1, 2, 3, 4, 7, 9, 10, 11, 12, 13), calcs=(5, 6, 8))
for r in range(6, 14):
    for c in range(3, 14):
        cp.cell(row=r, column=c).alignment = CTR
    cp.cell(row=r, column=6).number_format = "0.0%"
    cp.cell(row=r, column=8).number_format = "0.0"
si_no(cp, "M6:M13")

cp.cell(row=14, column=1, value="PROMEDIO PONDERADO")
cp.cell(row=14, column=3, value='=IF(SUM(C6:C13)=0,"",SUM(C6:C13))')
cp.cell(row=14, column=4, value='=IF(SUM(D6:D13)=0,"",SUM(D6:D13))')
cp.cell(row=14, column=5, value='=IF(SUM(E6:E13)=0,"",SUM(E6:E13))')
cp.cell(row=14, column=6, value='=IF(SUM($E$6:$E$13)=0,"",SUM($C$6:$C$13)/SUM($E$6:$E$13))')
cp.cell(row=14, column=7, value='=IF(SUM(G6:G13)=0,"",SUM(G6:G13))')
cp.cell(row=14, column=8, value='=IF(OR(SUM($E$6:$E$13)=0,SUM($G$6:$G$13)=0),"",SUM($E$6:$E$13)/SUM($G$6:$G$13))')
for c in (9, 10, 11, 12):
    L = get_column_letter(c)
    cp.cell(row=14, column=c, value='=IF(SUM({0}6:{0}13)=0,"",SUM({0}6:{0}13))'.format(L))
for c in range(1, 14):
    cell = cp.cell(row=14, column=c)
    cell.font = font(10, True); cell.fill = FILL_S; cell.border = BOX
    cell.alignment = CTR if c > 1 else Alignment(vertical="center")
cp.cell(row=14, column=6).number_format = "0.0%"
cp.cell(row=14, column=8).number_format = "0.0"
cp.row_dimensions[14].height = 22
cp.conditional_formatting.add("F14",
    CellIsRule(operator="greaterThan", formula=["0.5"], fill=PatternFill("solid", fgColor="F9E4E3"), font=font(10, True, "A32320")))

cp.merge_cells("A16:M17")
c = cp["A16"]
c.value = ("Lectura: si caminar pasa del 50% del ciclo, el problema es el slotting y no la gente — correr más rápido no "
           "arregla una ruta mal diseñada. Las columnas de incidencias (no encontró, posición vacía) apuntan a exactitud "
           "de inventario; si el operario se salta el escaneo, ningún control posterior va a sostener la calidad.")
c.font = font(9, color="74818D"); c.alignment = WRAP

# =========================================================
# 4 CONTEO CIEGO
# =========================================================
cc_ = wb.create_sheet("4 Conteo ciego")
cc_.sheet_view.showGridLines = False
title_block(cc_, "CONTEO CIEGO — EXACTITUD DE INVENTARIO (ERI)",
            "100 ubicaciones elegidas por ti, no por ellos. A ciegas: quien cuenta no ve la cantidad del sistema. La columna Cant. sistema se llena DESPUÉS de contar.", 7)

lab = [("A4", "Ubicaciones contadas"), ("C4", "Exactas"), ("E4", "ERI por ubicación")]
for ref, txt in lab:
    cc_[ref] = txt; cc_[ref].font = font(10, True); cc_[ref].alignment = Alignment(horizontal="right", vertical="center")
cc_["B4"] = '=COUNTIF($F$8:$F$107,"SÍ")+COUNTIF($F$8:$F$107,"NO")'
cc_["D4"] = '=COUNTIF($F$8:$F$107,"SÍ")'
cc_["F4"] = '=IF($B$4=0,"",$D$4/$B$4)'
for ref, nf in [("B4", "0"), ("D4", "0"), ("F4", "0.0%")]:
    cell = cc_[ref]; cell.number_format = nf; cell.fill = FILL_C
    cell.font = font(12, True, NAVY_D); cell.alignment = CTR; cell.border = BOX
cc_.row_dimensions[4].height = 24
cc_.conditional_formatting.add("F4",
    CellIsRule(operator="lessThan", formula=["0.95"], fill=PatternFill("solid", fgColor="F9E4E3"), font=font(12, True, "A32320")))

headers(cc_, 6, ["#", "Ubicación", "SKU", "Cant. contada", "Cant. sistema", "¿Exacta?", "Diferencia"],
        [6, 18, 18, 14, 14, 11, 12])
cc_.freeze_panes = "A8"

cc_.cell(row=7, column=1, value="EJ")
cc_.cell(row=7, column=2, value="A-12-03-B"); cc_.cell(row=7, column=3, value="SKU 100482")
cc_.cell(row=7, column=4, value=48); cc_.cell(row=7, column=5, value=52)
cc_.cell(row=7, column=6, value='=IF(OR($D7="",$E7=""),"",IF($D7=$E7,"SÍ","NO"))')
cc_.cell(row=7, column=7, value='=IF(OR($D7="",$E7=""),"",$D7-$E7)')
example_row(cc_, 7, 7)

for i, r in enumerate(range(8, 108), start=1):
    cc_.cell(row=r, column=1, value=i)
    cc_.cell(row=r, column=6, value='=IF(OR($D{0}="",$E{0}=""),"",IF($D{0}=$E{0},"SÍ","NO"))'.format(r))
    cc_.cell(row=r, column=7, value='=IF(OR($D{0}="",$E{0}=""),"",$D{0}-$E{0})'.format(r))
    cc_.row_dimensions[r].height = 16
style_range(cc_, 8, 107, 1, 7, inputs=(2, 3, 4, 5), calcs=(6, 7))
for r in range(8, 108):
    for c in (1, 4, 5, 6, 7):
        cc_.cell(row=r, column=c).alignment = CTR
    cc_.cell(row=r, column=1).fill = FILL_S
    cc_.cell(row=r, column=1).font = font(9, color="74818D")
cc_.conditional_formatting.add("F8:F107",
    CellIsRule(operator="equal", formula=['"NO"'], fill=PatternFill("solid", fgColor="F9E4E3"), font=font(10, True, "A32320")))

# =========================================================
# 5 MUESTREO
# =========================================================
ms = wb.create_sheet("5 Muestreo")
ms.sheet_view.showGridLines = False
title_block(ms, "MUESTREO DE TRABAJO",
            "Cada 20 minutos, desde un punto alto, anota qué hace cada persona que veas EN ESE INSTANTE. Sin juzgar, solo el instante. Diez a doce rondas dan una muestra válida.", 9)
headers(ms, 4, ["Ronda", "Hora", "Persona / puesto", "Actividad", "", "Actividad", "Observaciones", "%", ""],
        [8, 9, 22, 20, 3, 20, 14, 10, 3])
for col in ("E", "I"):
    for r_ in range(4, 6):
        ms["{}{}".format(col, r_)].fill = PatternFill("solid", fgColor="FFFFFF")
        ms["{}{}".format(col, r_)].border = Border()
ms["E4"].value = None; ms["I4"].value = None
ms.freeze_panes = "A6"

acts = ["Alistando", "Reabasteciendo", "Caminando", "Buscando", "Esperando", "Reprocesando", "Otro"]
ms.cell(row=5, column=1, value="EJ"); ms.cell(row=5, column=2, value="09:20")
ms.cell(row=5, column=3, value="Picking pasillo 8"); ms.cell(row=5, column=4, value="Caminando")
example_row(ms, 5, 4)

for r in range(6, 206):
    ms.row_dimensions[r].height = 15
style_range(ms, 6, 205, 1, 4, inputs=(1, 2, 3, 4))
for r in range(6, 206):
    for c in (1, 2, 4):
        ms.cell(row=r, column=c).alignment = CTR

dv = DataValidation(type="list", formula1='"{}"'.format(",".join(acts)), allow_blank=True)
ms.add_data_validation(dv); dv.add("D6:D205")

for i, a in enumerate(acts):
    r = 5 + i
    ms.cell(row=r, column=6, value=a).font = font(10)
    ms.cell(row=r, column=7, value='=COUNTIF($D$6:$D$205,$F{})'.format(r))
    ms.cell(row=r, column=8, value='=IF($G$12=0,"",$G{}/$G$12)'.format(r))
    for c in (6, 7, 8):
        cell = ms.cell(row=r, column=c); cell.border = BOX; cell.fill = FILL_C
        if c > 6: cell.alignment = CTR
    ms.cell(row=r, column=8).number_format = "0.0%"
ms.cell(row=12, column=6, value="TOTAL OBSERVACIONES").font = font(10, True)
ms.cell(row=12, column=7, value="=SUM($G$5:$G$11)")
ms.cell(row=12, column=8, value='=IF($G$12=0,"",1)')
for c in (6, 7, 8):
    cell = ms.cell(row=12, column=c); cell.border = BOX; cell.fill = FILL_S; cell.font = font(10, True)
    if c > 6: cell.alignment = CTR
ms.cell(row=12, column=8).number_format = "0.0%"

ms.cell(row=14, column=6, value="No agrega valor").font = font(10, True, "A32320")
ms.cell(row=14, column=7, value="=$G$7+$G$8+$G$9+$G$10+$G$11")
ms.cell(row=14, column=8, value='=IF($G$12=0,"",$G$14/$G$12)')
for c in (6, 7, 8):
    cell = ms.cell(row=14, column=c); cell.border = BOX; cell.fill = FILL_C; cell.font = font(11, True)
    if c > 6: cell.alignment = CTR
ms.cell(row=14, column=8).number_format = "0.0%"
ms.conditional_formatting.add("H14",
    CellIsRule(operator="greaterThan", formula=["0.5"], fill=PatternFill("solid", fgColor="F9E4E3"), font=font(11, True, "A32320")))

ms.merge_cells("F16:H20")
c = ms["F16"]
c.value = ("Referencia: en CEDI sin estándares de trabajo es normal encontrar 30–45% del tiempo en desplazamiento y "
           "10–20% buscando. Es el dato que convierte “la gente no rinde” en “el diseño no deja rendir”.")
c.font = font(9, color="74818D"); c.alignment = WRAP

# =========================================================
# 6 VEHÍCULOS
# =========================================================
vh = wb.create_sheet("6 Vehículos")
vh.sheet_view.showGridLines = False
title_block(vh, "PERMANENCIA DE VEHÍCULOS EN MUELLE",
            "Planilla de portería. Escribe las horas en formato HH:MM (por ejemplo 08:15). Los minutos se calculan solos.", 8)
headers(vh, 4, ["Placa", "Transportista", "Hora llegada", "Hora entrada muelle", "Hora salida",
                "Espera (min)", "Cargue (min)", "Permanencia total (min)"],
        [12, 20, 13, 17, 12, 13, 13, 20])
vh.freeze_panes = "A5"

vh.cell(row=5, column=1, value="EJ — ABC123"); vh.cell(row=5, column=2, value="Transportes Norte")
for c, t in [(3, "07:40"), (4, "09:05"), (5, "10:20")]:
    vh.cell(row=5, column=c, value=t)
vh.cell(row=5, column=6, value='=IF(OR($C5="",$D5=""),"",($D5-$C5)*1440)')
vh.cell(row=5, column=7, value='=IF(OR($D5="",$E5=""),"",($E5-$D5)*1440)')
vh.cell(row=5, column=8, value='=IF(OR($C5="",$E5=""),"",($E5-$C5)*1440)')
example_row(vh, 5, 8)

for r in range(6, 26):
    vh.cell(row=r, column=6, value='=IF(OR($C{0}="",$D{0}=""),"",($D{0}-$C{0})*1440)'.format(r))
    vh.cell(row=r, column=7, value='=IF(OR($D{0}="",$E{0}=""),"",($E{0}-$D{0})*1440)'.format(r))
    vh.cell(row=r, column=8, value='=IF(OR($C{0}="",$E{0}=""),"",($E{0}-$C{0})*1440)'.format(r))
    vh.row_dimensions[r].height = 17
style_range(vh, 6, 25, 1, 8, inputs=(1, 2, 3, 4, 5), calcs=(6, 7, 8))
for r in range(6, 26):
    for c in range(3, 9):
        vh.cell(row=r, column=c).alignment = CTR
        if c >= 6: vh.cell(row=r, column=c).number_format = "0"
        else: vh.cell(row=r, column=c).number_format = "hh:mm"

vh.cell(row=26, column=1, value="PROMEDIO")
for c in (6, 7, 8):
    L = get_column_letter(c)
    vh.cell(row=26, column=c, value='=IF(COUNT({0}6:{0}25)=0,"",AVERAGE({0}6:{0}25))'.format(L))
for c in range(1, 9):
    cell = vh.cell(row=26, column=c); cell.font = font(11, True); cell.fill = FILL_S; cell.border = BOX
    cell.alignment = CTR if c > 1 else Alignment(vertical="center")
    if c >= 6: cell.number_format = "0"
vh.row_dimensions[26].height = 22
vh.conditional_formatting.add("H26",
    CellIsRule(operator="greaterThan", formula=["90"], fill=PatternFill("solid", fgColor="F9E4E3"), font=font(11, True, "A32320")))

vh.merge_cells("A28:H29")
c = vh["A28"]
c.value = ("Lectura: si la espera pesa más que el cargue, el cuello está en programación de citas y no en la operación de muelle. "
           "Permanencias sobre tres horas suelen significar staging saturado con pedidos ya alistados esperando vehículo.")
c.font = font(9, color="74818D"); c.alignment = WRAP

# =========================================================
# 7 PALLETS EN PISO
# =========================================================
pp = wb.create_sheet("7 Pallets en piso")
pp.sheet_view.showGridLines = False
title_block(pp, "PALLETS EN PISO — PRUEBA DE LA ETIQUETA DE FECHA",
            "Etiqueta con fecha y hora cada pallet fuera de posición y cuéntalos por zona. Vuelve a las 48 horas: lo que siga ahí con la etiqueta original es el backlog real.", 6)
headers(pp, 4, ["Zona", "N.º de pallets etiquetados", "Producto / descripción",
                "Fecha de la etiqueta", "¿Sigue ahí a las 48 h?", "Observación"],
        [24, 20, 30, 16, 18, 30])
pp.freeze_panes = "A5"

pp.cell(row=5, column=1, value="EJ — Recepción"); pp.cell(row=5, column=2, value=18)
pp.cell(row=5, column=3, value="Importado sin ubicar, 3 referencias")
pp.cell(row=5, column=4, value="24/08/2026"); pp.cell(row=5, column=5, value="Sí")
example_row(pp, 5, 6)

zonas = ["Recepción", "Staging de despacho", "Pasillos de maniobra", "Devoluciones",
         "Averías", "Cuarentena / pendiente calidad", "Cross-dock", "Otro", "", ""]
for i, z in enumerate(zonas):
    r = 6 + i
    pp.cell(row=r, column=1, value=z if z else None)
    pp.row_dimensions[r].height = 20
style_range(pp, 6, 15, 1, 6, inputs=(1, 2, 3, 4, 5, 6))
for r in range(6, 16):
    pp.cell(row=r, column=2).alignment = CTR
    pp.cell(row=r, column=4).alignment = CTR
    pp.cell(row=r, column=5).alignment = CTR
si_no(pp, "E6:E15")

pp.cell(row=16, column=1, value="TOTAL PALLETS EN PISO")
pp.cell(row=16, column=2, value="=SUM(B6:B15)")
for c in range(1, 7):
    cell = pp.cell(row=16, column=c); cell.font = font(11, True); cell.fill = FILL_S; cell.border = BOX
pp.cell(row=16, column=2).alignment = CTR
pp.row_dimensions[16].height = 22

# =========================================================
# 8 AGOTADOS
# =========================================================
ag = wb.create_sheet("8 Agotados")
ag.sheet_view.showGridLines = False
title_block(ag, "AGOTADOS INTERNOS DE PICKING",
            "Cada vez que una posición de picking queda vacía con pedido pendiente. Déjalo con un supervisor para todo el turno: es la causa oculta número uno de baja productividad y casi nadie la mide.", 6)
headers(ag, 4, ["Hora", "SKU", "Ubicación", "Minutos sin reabastecer", "Pedido afectado", "Turno"],
        [10, 18, 16, 22, 18, 12])
ag.freeze_panes = "A6"

ag.cell(row=5, column=1, value="10:35"); ag.cell(row=5, column=2, value="SKU 100482")
ag.cell(row=5, column=3, value="A-04-01-A"); ag.cell(row=5, column=4, value=22)
ag.cell(row=5, column=5, value="PED-88213"); ag.cell(row=5, column=6, value="Mañana")
example_row(ag, 5, 6)

for r in range(6, 46):
    ag.row_dimensions[r].height = 16
style_range(ag, 6, 45, 1, 6, inputs=(1, 2, 3, 4, 5, 6))
for r in range(6, 46):
    for c in (1, 4, 6):
        ag.cell(row=r, column=c).alignment = CTR

ag.cell(row=46, column=1, value="EVENTOS")
ag.cell(row=46, column=2, value="=COUNTA(B6:B45)")
ag.cell(row=46, column=3, value="MINUTOS PERDIDOS")
ag.cell(row=46, column=4, value='=IF(SUM(D6:D45)=0,"",SUM(D6:D45))')
for c in range(1, 7):
    cell = ag.cell(row=46, column=c); cell.font = font(11, True); cell.fill = FILL_S; cell.border = BOX
ag.cell(row=46, column=2).alignment = CTR
ag.cell(row=46, column=4).alignment = CTR
ag.row_dimensions[46].height = 22

# =========================================================
# 9 PREGUNTAS
# =========================================================
pr = wb.create_sheet("9 Preguntas")
pr.sheet_view.showGridLines = False
title_block(pr, "BANCO DE PREGUNTAS",
            "A los operarios, sin el supervisor delante. No preguntes “¿por qué está tan lleno?”: te dan la respuesta ensayada.", 4)
headers(pr, 4, ["Interlocutor", "Pregunta", "Respuesta", "Qué revela"], [18, 46, 46, 40])
pr.freeze_panes = "A5"

preguntas = [
    ("Operario", "¿Qué es lo que más tiempo le hace perder en el día?", "Nombra el desperdicio real, casi siempre antes que cualquier indicador."),
    ("Operario", "¿Cuántas veces al día no encuentra el producto donde dice el sistema?", "Estima el ERI sin esperar el conteo."),
    ("Operario", "Cuando se equivoca, ¿por qué cree que pasó?", "Distingue ambigüedad de diseño frente a falta de método."),
    ("Operario", "¿Cuántas líneas se supone que debe hacer por hora?", "Si no lo sabe, no hay estándar. Sin estándar no hay productividad que gestionar."),
    ("Operario", "¿Qué hay ahí que no debería estar?", "Señala inventario muerto y zonas sin dueño."),
    ("Jefe del CEDI", "¿Cuál es tu cuello de botella hoy?", "Si responde rápido y con un número, hay gestión. Si responde “todo”, no la hay."),
    ("Jefe del CEDI", "¿Cuántas posiciones tienes y cuántas están ocupadas en este momento?", "Contrasta contra tu conteo de la hoja 2."),
    ("Jefe del CEDI", "¿Qué pasa cuando un pedido no alcanza a salir?", "Revela si hay regla o improvisación."),
    ("Jefe del CEDI", "¿Cuándo fue el último re-slotting?", "Más de seis meses explica buena parte del desplazamiento."),
    ("Jefe del CEDI", "¿Cuánto se demora un camión desde que llega hasta que el producto queda disponible?", "Contrasta contra el dock-to-stock real."),
    ("Recepción", "¿Se recibe en flujo o por lotes al final del turno?", "El lote al cierre es lo que deja producto en piso toda la noche."),
    ("Despacho", "¿Qué los detiene más: esperar producto, esperar documento o esperar camión?", "Ubica el cuello sin necesidad de datos."),
    ("Calidad / devoluciones", "¿Quién autoriza la disposición final y cuándo fue la última vez?", "Zona sin dueño ni fecha límite: ahí se acumula el espacio."),
    ("Cualquiera", "Muéstrame el pallet más viejo que hay en piso.", "Sustituye a la pregunta “¿por qué está tan lleno?”."),
    ("Cualquiera", "Muéstrame diez posiciones que no se hayan movido este mes.", "Caminando hacia allá ves más que con cualquier respuesta."),
]
r = 5
for who, q, rev in preguntas:
    pr.cell(row=r, column=1, value=who).font = font(9, True, TAPE)
    pr.cell(row=r, column=2, value=q)
    pr.cell(row=r, column=4, value=rev).font = font(9, color="74818D")
    pr.row_dimensions[r].height = 30
    r += 1
style_range(pr, 5, r - 1, 1, 4, inputs=(3,))
for rr in range(5, r):
    pr.cell(row=rr, column=1).font = font(9, True, TAPE)
    pr.cell(row=rr, column=4).font = font(9, color="74818D")

# =========================================================
# 10 SOLICITUD DATOS
# =========================================================
sd = wb.create_sheet("10 Solicitud datos")
sd.sheet_view.showGridLines = False
title_block(sd, "SOLICITUD DE INFORMACIÓN",
            "Entrégala hoy, por escrito, con responsable y fecha. Período: últimas 13 semanas. Formato Excel o CSV plano, un registro por fila — sin tablas dinámicas ni consolidados.", 7)
sd["E3"] = "Avance de entrega"
sd["E3"].font = font(10, True); sd["E3"].alignment = Alignment(horizontal="right", vertical="center")
sd["G3"] = '=IF(COUNTA($C$5:$C$20)=0,"",COUNTIF($G$5:$G$20,"Sí")/COUNTA($C$5:$C$20))'
sd["G3"].number_format = "0.0%"; sd["G3"].font = font(11, True, NAVY_D)
sd["G3"].fill = FILL_C; sd["G3"].alignment = CTR; sd["G3"].border = BOX
headers(sd, 4, ["#", "Área", "Información solicitada", "Formato", "Responsable", "Fecha compromiso", "Recibido"],
        [5, 18, 58, 16, 20, 16, 11])
sd.freeze_panes = "A5"

datos = [
    ("WMS / ERP", "Movimientos de salida a nivel de línea: fecha, hora, SKU, cantidad, ubicación, operario, pedido, cliente y ruta", "CSV plano"),
    ("WMS / ERP", "Movimientos de entrada con hora de llegada del vehículo Y hora de ubicación del producto (dock-to-stock)", "CSV plano"),
    ("WMS / ERP", "Snapshot de inventario por ubicación, al cierre de hoy", "CSV plano"),
    ("WMS / ERP", "Maestro de ubicaciones: posiciones por zona y por tipo", "Excel"),
    ("WMS / ERP", "Maestro de SKU: dimensiones, peso, empaque y paletización", "Excel"),
    ("WMS / ERP", "Trazabilidad de pedidos: creación, corte, liberación, fin de alistamiento y despacho", "CSV plano"),
    ("WMS / ERP", "Devoluciones y notas crédito con su causal", "CSV plano"),
    ("WMS / ERP", "Resultados de los últimos conteos cíclicos", "Excel"),
    ("Gestión humana", "Headcount por turno y por función", "Excel"),
    ("Gestión humana", "Horas ordinarias y horas extra por período", "Excel"),
    ("Gestión humana", "Ausentismo, rotación y antigüedad promedio", "Excel"),
    ("Transporte", "Vehículos programados frente a ejecutados", "Excel"),
    ("Transporte", "Hora de llegada y de salida por vehículo", "CSV plano"),
    ("Transporte", "Ocupación del vehículo despachado", "Excel"),
    ("Sitio", "Layout a escala con zonas, muelles y pasillos", "PDF / DWG"),
    ("Finanzas", "Costo total del CEDI del último trimestre", "Excel"),
]
r = 5
for i, (area, info, fmt) in enumerate(datos, start=1):
    sd.cell(row=r, column=1, value=i)
    sd.cell(row=r, column=2, value=area)
    sd.cell(row=r, column=3, value=info)
    sd.cell(row=r, column=4, value=fmt)
    sd.row_dimensions[r].height = 28
    r += 1
style_range(sd, 5, 20, 1, 7, inputs=(5, 6, 7))
for rr in range(5, 21):
    sd.cell(row=rr, column=1).alignment = CTR
    sd.cell(row=rr, column=2).font = font(9, True, TAPE)
    sd.cell(row=rr, column=4).alignment = CTR
    sd.cell(row=rr, column=4).font = font(9, color="74818D")
    sd.cell(row=rr, column=7).alignment = CTR
si_no(sd, "G5:G20")
sd.conditional_formatting.add("G5:G20",
    CellIsRule(operator="equal", formula=['"Sí"'], fill=PatternFill("solid", fgColor="DFEFE7")))

sd.merge_cells("A22:G24")
c = sd["A22"]
c.value = ("Si te mandan un reporte ya cocinado, no sirve: necesitas el detalle línea por línea. Un consolidado entrega "
           "promedios, y los promedios son precisamente donde se esconden el pico de las últimas tres horas antes del corte, "
           "el operario con menos de 90 días y los diez SKU que generan la mitad de los errores.")
c.font = font(9, color="74818D"); c.alignment = WRAP

wb.save(OUT)
print("OK ->", OUT)
